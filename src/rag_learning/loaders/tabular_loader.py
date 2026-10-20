"""CSV / Excel 解析：每个表（工作表）转成 Markdown 表格，保留行列结构。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd

from ..schema import ParsedDocument
from .base import AclOverride, build_document, dataframe_to_markdown, rows_to_markdown_table


def _load_csv(path: Path) -> str:
    """读取 CSV 并转 Markdown，编码依次尝试 UTF-8-BOM 和 GB18030。"""
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            # dtype=str 保留原始写法，避免 "007" 被解析成 7；
            # keep_default_na=False 让空单元格保持空串而不是 NaN
            frame = pd.read_csv(path, dtype=str, keep_default_na=False, encoding=encoding)
            return f"# {path.stem}\n\n{dataframe_to_markdown(frame)}"
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(f"CSV 编码无法识别：{last_error}")


def _load_xlsx(path: Path) -> str:
    """逐工作表读取，每个 sheet 输出一个二级标题加一张 Markdown 表格。"""
    from openpyxl import load_workbook

    # read_only + data_only：流式读取、只取公式计算后的值
    workbook = load_workbook(path, data_only=True, read_only=True)
    sections: list[str] = [f"# {path.stem}"]
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                # 整行空白的工作表区域直接跳过，减少噪声
                if row and any(cell not in (None, "") for cell in row):
                    rows.append(list(row))
            if not rows:
                continue
            sections.append(f"## 工作表：{sheet.title}")
            sections.append(rows_to_markdown_table(rows))
    finally:
        workbook.close()
    return "\n\n".join(sections)


def load(path: Path, docs_root: Path, override: AclOverride | None = None) -> ParsedDocument:
    """按扩展名分发 CSV / XLSX，统一交给 build_document 收尾。"""
    suffix = path.suffix.lower()
    if suffix == ".csv":
        text = _load_csv(path)
        doc_format = "csv"
    else:
        text = _load_xlsx(path)
        doc_format = "xlsx"
    return build_document(path, docs_root, text, doc_format, {}, override)
