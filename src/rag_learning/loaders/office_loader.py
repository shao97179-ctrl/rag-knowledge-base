"""DOCX / PPTX 解析：按文档顺序保留标题、正文、表格和幻灯片页码。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from ..schema import ParsedDocument
from .base import AclOverride, build_document, rows_to_markdown_table


def _docx_text(path: Path) -> str:
    """按 body 里的原始顺序遍历段落和表格，把样式映射成 Markdown 标记。"""
    document = Document(str(path))
    sections: list[str] = []
    # 直接遍历 XML 子节点，而不是 document.paragraphs / tables，
    # 这样表格和段落的相对顺序不会被打乱
    for child in document.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            style = (paragraph.style.name or "") if paragraph.style is not None else ""
            # Heading 1/2/3 -> #/##/###，最多 6 级
            if style.lower().startswith("heading"):
                digits = "".join(char for char in style if char.isdigit())
                level = min(int(digits), 6) if digits else 1
                sections.append(f"{'#' * level} {text}")
            elif style.lower() == "title":
                sections.append(f"# {text}")
            elif "list" in style.lower():
                sections.append(f"- {text}")
            else:
                sections.append(text)
        elif child.tag == qn("w:tbl"):
            table = Table(child, document)
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            markdown = rows_to_markdown_table(rows)
            if markdown:
                sections.append(markdown)
    return "\n\n".join(sections)


def _pptx_text(path: Path) -> tuple[str, str | None]:
    """遍历每页幻灯片，输出页码标题、文本、表格和备注；同时返回首页标题。"""
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections: list[str] = []
    first_title: str | None = None
    for page_number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None and title_shape.text else f"第 {page_number} 页"
        if first_title is None and title_shape is not None and title_shape.text:
            first_title = title
        # 每页用一个二级标题分隔，chunker 会按页维护 heading 路径
        sections.append(f"## 第 {page_number} 页：{title}")

        def walk(shapes) -> None:
            """递归遍历 shape；组合形状要钻进去，否则会漏内容。"""
            for shape in shapes:
                # 标题已经作为分页标题输出过，这里跳过避免重复
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if shape.shape_type == 6:  # GROUP
                    walk(shape.shapes)
                    continue
                if getattr(shape, "has_table", False) and shape.has_table:
                    rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                    markdown = rows_to_markdown_table(rows)
                    if markdown:
                        sections.append(markdown)
                    continue
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        # 多级项目符号用缩进层级近似成列表项
                        sections.append(f"- {text}" if paragraph.level > 0 else text)

        walk(slide.shapes)
        # 备注里常有讲稿补充信息，对问答有价值，一并入库
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                sections.append(f"### 备注\n{notes}")
    return "\n\n".join(sections), first_title


def load(path: Path, docs_root: Path, override: AclOverride | None = None) -> ParsedDocument:
    """按扩展名分发到 DOCX / PPTX 解析，再走公共的清洗和权限推断。"""
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text = _docx_text(path)
        doc_format = "docx"
        frontmatter = {}
    elif suffix == ".pptx":
        text, first_title = _pptx_text(path)
        doc_format = "pptx"
        frontmatter = {"title": first_title} if first_title else {}
    else:
        raise ValueError(f"不支持的 Office 格式：{suffix}")
    return build_document(path, docs_root, text, doc_format, frontmatter, override)
